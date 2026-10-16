from pathlib import Path
from openpyxl import load_workbook
import re
from pypdf import PdfReader
import json

root=Path(__file__).parent
w=load_workbook(root/'Gatto_Chepeliev_Supplementary_info_database.xlsx',read_only=True,data_only=True)
print('README',list(w['READ ME'].values))
for name in ['Magnitude','Nutrition']:
    print(name,list(w[name].iter_rows(min_row=1,max_row=5,values_only=True)))
html=(root/'WRAP_Quelle.html').read_text(encoding='utf8')
print('WRAP PDFs',re.findall(r'href="([^"]+\.pdf[^"]*)"',html))
r=PdfReader(root/'REFOWAS_2021.pdf')
(root/'REFOWAS_2021.txt').write_text('\n'.join('PAGE '+str(i+1)+'\n'+p.extract_text() for i,p in enumerate(r.pages)),encoding='utf8')
for n in ['WRAP_Households_2022_2025','WRAP_Methods_2022_2025','FAO_FBS_Guidelines','Friedman_Heiman_2024','UNEP_FWI_2024']:
    try:
        r=PdfReader(root/(n+'.pdf'))
        (root/(n+'.txt')).write_text('\n'.join('PAGE '+str(i+1)+'\n'+p.extract_text() for i,p in enumerate(r.pages)),encoding='utf8')
        print(n,len(r.pages))
    except Exception as e:
        print(n,type(e).__name__)
for name in ['Magnitude','Nutrition']:
    s=w[name]
    print(name,'years',sorted(set(row[0] for row in s.iter_rows(min_row=2,values_only=True) if row[0] is not None)))
